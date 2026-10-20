"""Sampling and agreement analysis for the human semantic-validation pilot.

Task A uses one canonical four-label schema; see
docs/task_a_semantic_pilot_guidelines.md.
"""

from collections import Counter, defaultdict
import csv
import itertools
import json
import math
from pathlib import Path
import random

from task_a.leakage import comparison_key

# Canonical Task A schema. Task B's SUPPORTED/CONTRADICTED/UNSUPPORTED labels are unrelated.
LABELS = ("BENIGN", "HARMFUL", "ADVERSARIAL", "QUARANTINE")
REVIEWER_COUNT = 3
EXPECTED_ITEMS = 200
REQUIRED_COLUMNS = ("pilot_id", "prompt", "safellm_label", "reviewer_id", "rationale")


def sample_unique_strata(rows, quota=50, seed=42):
    """Sample comparison-normalized unique prompts from four source strata."""
    if quota < 1:
        raise ValueError("quota must be positive")
    strata = defaultdict(dict)
    duplicate_candidates = Counter()
    for index, row in enumerate(rows):
        prompt = row.get("prompt")
        flag = row.get("adversarial")
        harm = row.get("prompt_harm_label")
        if not isinstance(prompt, str) or not prompt.strip():
            continue
        if type(flag) is not bool or harm not in {"harmful", "unharmful"}:
            continue
        stratum = f"{str(flag).lower()}__{harm}"
        key = comparison_key(prompt)
        duplicate_candidates[stratum] += key in strata[stratum]
        strata[stratum].setdefault(key, (index, row))

    selected = []
    summary = {}
    for stratum in ("false__harmful", "false__unharmful",
                    "true__harmful", "true__unharmful"):
        candidates = list(strata[stratum].values())
        rng = random.Random(f"{seed}:{stratum}")
        chosen = rng.sample(candidates, min(quota, len(candidates)))
        summary[stratum] = {
            "eligible_unique": len(candidates),
            "selected": len(chosen),
            "shortfall": max(0, quota - len(chosen)),
            "duplicate_candidates_skipped": duplicate_candidates[stratum],
        }
        for index, row in chosen:
            selected.append({
                "pilot_id": f"pilot-{len(selected) + 1:03d}",
                "source_id": f"wildguard:train:{index}",
                "prompt": row["prompt"],
                "source_adversarial": row["adversarial"],
                "source_harmfulness": row["prompt_harm_label"],
                "source_risk_category": row.get("subcategory"),
                "source_split": "train",
                "sampling_stratum": stratum,
            })
    return selected, summary


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _xlsx_records(path):
    from openpyxl import load_workbook  # needed only for spreadsheet inputs

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []
        names = ["" if name is None else str(name) for name in header]
        return [{name: _cell_text(value) for name, value in zip(names, row)}
                for row in rows if any(value is not None for value in row)]
    finally:
        workbook.close()


def read_annotations(path, expected_reviewer_id, expected_ids=None):
    """Validate one completed reviewer file (.csv or .xlsx) without modifying it.

    Returns rows in file order as dicts with pilot_id, prompt, label and rationale.
    """
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        rows = _xlsx_records(path)
    else:
        with open(path, newline="", encoding="utf-8-sig") as source:
            rows = list(csv.DictReader(source))
    if not rows or not set(REQUIRED_COLUMNS).issubset(rows[0]):
        raise ValueError(f"{path}: missing required columns or rows")
    found = [row["pilot_id"] for row in rows]
    if len(found) != len(set(found)):
        raise ValueError(f"{path}: duplicate pilot IDs")
    if expected_ids is not None and found != list(expected_ids):
        raise ValueError(f"{path}: IDs must match the expected pilot IDs and order")
    result = []
    for row in rows:
        label = (row["safellm_label"] or "").strip().upper()
        if label not in LABELS:
            raise ValueError(f'{path}: invalid label {row["safellm_label"]!r} '
                             f'for {row["pilot_id"]}; allowed: {", ".join(LABELS)}')
        if (row["reviewer_id"] or "").strip() != str(expected_reviewer_id):
            raise ValueError(f"{path}: reviewer_id must be {expected_reviewer_id}")
        rationale = (row["rationale"] or "").strip()
        if not rationale:
            raise ValueError(f'{path}: rationale required for {row["pilot_id"]}')
        result.append({"pilot_id": row["pilot_id"], "prompt": row["prompt"],
                       "label": label, "rationale": rationale})
    return result


def expected_pilot_ids(items=EXPECTED_ITEMS):
    """The pilot's ID set and order, as assigned by sample_unique_strata."""
    return [f"pilot-{i:03d}" for i in range(1, items + 1)]


def load_reviews(files, expected_ids=None):
    """Read reviewer files {reviewer_id: path}.

    Requires the exact expected pilot ID list (default: pilot-001..pilot-200, in order),
    so identically truncated or reordered files cannot pass, plus identical prompts.
    """
    if len(files) != REVIEWER_COUNT:
        raise ValueError(f"Exactly {REVIEWER_COUNT} reviewer files are required")
    expected_ids = expected_pilot_ids() if expected_ids is None else list(expected_ids)
    reviews = {reviewer: read_annotations(path, reviewer, expected_ids)
               for reviewer, path in files.items()}
    reference_id, reference = next(iter(reviews.items()))
    for reviewer, rows in reviews.items():
        if [r["pilot_id"] for r in rows] != [r["pilot_id"] for r in reference]:
            raise ValueError(f"Reviewer {reviewer} IDs/order differ from reviewer {reference_id}")
        if [r["prompt"] for r in rows] != [r["prompt"] for r in reference]:
            raise ValueError(f"Reviewer {reviewer} prompt text differs from reviewer {reference_id}")
    return reviews


def _rate(count, total):
    return count / total if total else None


def cohens_kappa(labels_a, labels_b):
    if len(labels_a) != len(labels_b) or not labels_a:
        raise ValueError("Two nonempty equal-length label lists are required")
    n = len(labels_a)
    observed = sum(a == b for a, b in zip(labels_a, labels_b)) / n
    counts_a, counts_b = Counter(labels_a), Counter(labels_b)
    expected = sum((counts_a[label] / n) * (counts_b[label] / n) for label in LABELS)
    if math.isclose(expected, 1.0):
        return 1.0 if math.isclose(observed, 1.0) else None
    return (observed - expected) / (1 - expected)


def fleiss_kappa(label_lists):
    """Fleiss' kappa for equal-length label lists, one per rater (two or more)."""
    raters = len(label_lists)
    if raters < 2 or not label_lists[0] or len({len(x) for x in label_lists}) != 1:
        raise ValueError("At least two nonempty equal-length label lists are required")
    items = len(label_lists[0])
    totals = Counter()
    observed = 0.0
    for item in zip(*label_lists):
        counts = Counter(item)
        totals.update(counts)
        observed += sum(c * (c - 1) for c in counts.values()) / (raters * (raters - 1))
    observed /= items
    expected = sum((totals[label] / (items * raters)) ** 2 for label in LABELS)
    if math.isclose(expected, 1.0):
        return 1.0 if math.isclose(observed, 1.0) else None
    return (observed - expected) / (1 - expected)


def confusion_matrix(labels_a, labels_b):
    """matrix[label_from_first][label_from_second] -> count, over all four labels."""
    matrix = {a: {b: 0 for b in LABELS} for a in LABELS}
    for a, b in zip(labels_a, labels_b):
        matrix[a][b] += 1
    return matrix


def agreement_pattern(labels):
    return {1: "unanimous", 2: "two_vs_one"}.get(len(set(labels)), "three_way")


def disagreement_category(labels_by_reviewer):
    """Descriptive only; says nothing about which reviewer is right."""
    counts = Counter(labels_by_reviewer.values())
    if len(counts) == 1:
        return "unanimous"
    label, size = counts.most_common(1)[0]
    if size * 2 > len(labels_by_reviewer):
        dissent = ", ".join(f"{r}={l}" for r, l in labels_by_reviewer.items() if l != label)
        return f"majority {label} ({size}/{len(labels_by_reviewer)}); dissent: {dissent}"
    return f"no majority ({len(counts)} distinct labels)"


def stratum_disagreement(master, reviews):
    """Items with any disagreement per source sampling stratum (needs the local master)."""
    by_id = {row["pilot_id"]: row for row in master}
    first = next(iter(reviews.values()))
    if set(by_id) != {row["pilot_id"] for row in first}:
        raise ValueError("Master pilot IDs must match reviewer pilot IDs")
    strata = defaultdict(lambda: {"items": 0, "items_with_any_disagreement": 0})
    for i, row in enumerate(first):
        stat = strata[by_id[row["pilot_id"]]["sampling_stratum"]]
        stat["items"] += 1
        stat["items_with_any_disagreement"] += len({r[i]["label"] for r in reviews.values()}) > 1
    return dict(strata)


def analyze_agreement(reviews, include_prompts=False):
    """Observed agreement among three submitted annotation sets, four labels kept separate.

    Measures agreement among the submitted annotations only. Reviewer independence is
    owner-confirmed, not established or audited here; this also does not establish
    correctness or agreement with any source-dataset label.
    """
    ids = list(reviews)
    if len(ids) != REVIEWER_COUNT:
        raise ValueError(f"Exactly {REVIEWER_COUNT} reviewers are required")
    n = len(reviews[ids[0]])
    labels = {r: [row["label"] for row in reviews[r]] for r in ids}

    distributions = {}
    for r in ids:
        counts = Counter(labels[r])
        distributions[r] = {
            "counts": {label: counts[label] for label in LABELS},
            "quarantine_count": counts["QUARANTINE"],
            "quarantine_rate": _rate(counts["QUARANTINE"], n),
        }

    pairwise = {}
    for a, b in itertools.combinations(ids, 2):
        agree = sum(x == y for x, y in zip(labels[a], labels[b]))
        pairwise[f"{a}-{b}"] = {
            "raw_agreement_count": agree,
            "raw_agreement_rate": _rate(agree, n),
            "cohens_kappa": cohens_kappa(labels[a], labels[b]),
            "confusion_matrix": confusion_matrix(labels[a], labels[b]),
            "confusion_matrix_axes": {"rows": f"reviewer {a}", "columns": f"reviewer {b}"},
        }

    patterns = {"unanimous": [], "two_vs_one": [], "three_way": []}
    disagreements = []
    for i in range(n):
        pattern = agreement_pattern([labels[r][i] for r in ids])
        pilot_id = reviews[ids[0]][i]["pilot_id"]
        patterns[pattern].append(pilot_id)
        if pattern == "unanimous":
            continue
        record = {"pilot_id": pilot_id, "pattern": pattern,
                  "category": disagreement_category({r: labels[r][i] for r in ids})}
        if include_prompts:
            record["prompt"] = reviews[ids[0]][i]["prompt"]
        for r in ids:
            record[f"reviewer_{r}_label"] = labels[r][i]
        for r in ids:
            record[f"reviewer_{r}_rationale"] = reviews[r][i]["rationale"]
        disagreements.append(record)

    return {
        "label_schema": list(LABELS),
        "reviewers": ids,
        "items": n,
        "label_distributions": distributions,
        "pairwise": pairwise,
        "fleiss_kappa": fleiss_kappa([labels[r] for r in ids]),
        "agreement_patterns": {
            k: {"count": len(v), "rate": _rate(len(v), n), "pilot_ids": v}
            for k, v in patterns.items()},
        "disagreement_rows": len(disagreements),
        "independence_verified": False,
        "independence_owner_confirmed": True,
    }, disagreements


def source_label_crosstab(master, reviews):
    """Descriptive source-flag x human-label counts, separate from agreement.

    No mapping from the four SafeLLM labels to the source boolean is applied, no
    mismatch rate is computed, and the source flag is not treated as ground truth.
    See docs/task_a_source_label_comparison.md.
    """
    source = {row["pilot_id"]: row for row in master}
    result = {}
    for reviewer, rows in reviews.items():
        if set(source) != {row["pilot_id"] for row in rows}:
            raise ValueError("Master pilot IDs must match reviewer pilot IDs")
        table = {"true": Counter(), "false": Counter()}
        for row in rows:
            table[str(source[row["pilot_id"]]["source_adversarial"]).lower()][row["label"]] += 1
        result[reviewer] = {flag: {label: counts[label] for label in LABELS}
                            for flag, counts in table.items()}
    return result


def render_report(results):
    """Markdown report for analyze_agreement() results."""
    ids = results["reviewers"]
    schema = results["label_schema"]

    def fmt(x):
        return "n/a" if x is None else f"{x:.4f}"

    lines = [
        "# Task A pilot: observed agreement among submitted annotations", "",
        f"Label schema (canonical, four labels kept separate): {', '.join(schema)}.",
        f"Items: {results['items']}. Reviewers: {', '.join(ids)}.", "",
        "**Limitation.** Reviewer independence is owner-confirmed, not independently "
        "audited: the project owner confirms the three reviewers annotated "
        "independently without seeing or influencing one another's annotations, and "
        "the repository cannot verify this. Reviewer 1's rationales contain "
        "review-dialogue language, Reviewer 2's rationales are repeated generic "
        "sentences, and file timestamps alone do not establish independence. These "
        "figures describe observed agreement among the submitted annotations and are "
        "not agreement with any source-dataset label.", ""]
    inputs = results.get("inputs")
    if inputs:
        lines += [
            "**Inputs.** The reviewer annotation files are " + inputs["location"] + ". These "
            "results cannot be reproduced from the repository alone; the files analysed are "
            "identified by SHA-256:", ""]
        lines += [f"- Reviewer {r}: `{f['name']}` `{f['sha256']}`"
                  for r, f in inputs["files"].items()]
        lines.append("")
    lines += [
        "## Label distributions", "",
        "| Reviewer | " + " | ".join(schema) + " | QUARANTINE rate |",
        "|---|" + "---|" * (len(schema) + 1)]
    for r in ids:
        d = results["label_distributions"][r]
        lines.append(f"| {r} | " + " | ".join(str(d["counts"][l]) for l in schema)
                     + f" | {d['quarantine_rate']:.1%} |")
    lines += ["", "## Pairwise agreement", "",
              "| Pair | Raw agreement | Rate | Cohen's kappa |", "|---|---|---|---|"]
    for pair, p in results["pairwise"].items():
        lines.append(f"| {pair} | {p['raw_agreement_count']}/{results['items']} | "
                     f"{p['raw_agreement_rate']:.4f} | {fmt(p['cohens_kappa'])} |")
    lines += ["", f"Fleiss' kappa (3 reviewers): **{fmt(results['fleiss_kappa'])}**", "",
              "## Agreement patterns", "", "| Pattern | Items | Rate |", "|---|---|---|"]
    for k, v in results["agreement_patterns"].items():
        lines.append(f"| {k} | {v['count']} | {v['rate']:.4f} |")
    for pair, p in results["pairwise"].items():
        a, b = pair.split("-")
        lines += ["", f"## Confusion matrix: reviewer {a} (rows) vs reviewer {b} (columns)", "",
                  "| | " + " | ".join(schema) + " |", "|---|" + "---|" * len(schema)]
        for row_label in schema:
            lines.append(f"| {row_label} | " + " | ".join(
                str(p["confusion_matrix"][row_label][c]) for c in schema) + " |")
    lines += ["", f"Per-item disagreements ({results['disagreement_rows']} rows) are in the "
              "accompanying disagreement table; none have been adjudicated.", ""]
    return "\n".join(lines)


def write_csv(path, fieldnames, rows):
    with open(path, "x", newline="", encoding="utf-8") as output:
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def jsonl_read(path):
    with open(path, encoding="utf-8") as source:
        return [json.loads(line) for line in source if line.strip()]
